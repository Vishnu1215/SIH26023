"""
DOCX Report Generator for Phase 8 using python-docx.
Produces formatted Microsoft Word (.docx) documents with official Ministry of Coal branding,
structured tables, metric summaries, and validation rule audit breakdowns.
"""

import os
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from app.services.report_utils import BRANDING, format_number, format_pct, format_datetime, clean_text


def _set_cell_background(cell, hex_color: str):
    """Set cell background color via XML manipulation."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def generate_docx_report(context: Dict[str, Any], output_path: str) -> str:
    """Generate official DOCX report and save to output_path."""
    doc = Document()

    # Set page margins to 0.7 inches
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

        # Header and Footer
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = header_p.add_run(f"MINISTRY OF COAL | CMPDI REPORTING PLATFORM  •  OFFICIAL USE")
        h_run.font.size = Pt(8)
        h_run.font.color.rgb = RGBColor(100, 116, 139)

        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        f_run = footer_p.add_run(f"SIH26023 Phase 8 Deterministic Engine  •  Generated: {context.get('formattedDate')}  •  Confidential")
        f_run.font.size = Pt(8)
        f_run.font.color.rgb = RGBColor(100, 116, 139)

    report_type = context.get("reportType", "executive").lower().replace(" ", "_").replace("-", "_")
    type_titles = {
        "executive": "Executive Mining & Analytics Report",
        "production": "National Coal Production & Subsidiary Performance Report",
        "validation": "Deterministic Validation & Data Discrepancy Audit",
        "dashboard": "Executive Dashboard Comprehensive Snapshot",
        "mine_performance": "Mine Performance & Extraction Register",
        "custom": "Custom Analytical & Compliance Report"
    }
    title = type_titles.get(report_type, "Ministry of Coal Statutory Report")

    # Ministry Branding Block
    p_rep = doc.add_paragraph()
    r_rep = p_rep.add_run(BRANDING["republic"])
    r_rep.font.size = Pt(9)
    r_rep.font.bold = True
    r_rep.font.color.rgb = RGBColor(100, 116, 139)

    p_min = doc.add_paragraph()
    r_min = p_min.add_run(BRANDING["ministry"])
    r_min.font.size = Pt(14)
    r_min.font.bold = True
    r_min.font.color.rgb = RGBColor(30, 58, 138)

    p_inst = doc.add_paragraph()
    r_inst = p_inst.add_run(BRANDING["institution"])
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(15, 23, 42)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(14)
    r_title = p_title.add_run(title)
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # Metadata Callout
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)
    m_run = p_meta.add_run(f"Generated: {context['formattedDate']}  |  Analytics Version: v{context['sourceAnalyticsVersion']}  |  Records Analyzed: {context['totalRecordsCount']}")
    m_run.font.size = Pt(9)
    m_run.font.italic = True
    m_run.font.color.rgb = RGBColor(71, 85, 105)

    dashboard = context.get("dashboard", {})
    prod = dashboard.get("production", {})
    val = dashboard.get("validation", {})
    docs = dashboard.get("documents", {})
    subs = dashboard.get("subsidiaries", [])
    mines = context.get("minesList", [])
    rule_violations = context.get("ruleViolations", [])

    # 1. Executive Summary Table (KPI Box)
    h_sec1 = doc.add_heading("1. Executive Overview & Key Production Metrics", level=2)
    h_sec1.paragraph_format.space_before = Pt(12)
    h_sec1.paragraph_format.space_after = Pt(6)

    table_kpi = doc.add_table(rows=3, cols=4)
    table_kpi.alignment = WD_TABLE_ALIGNMENT.CENTER
    kpi_headers = ["Total Coal Production", "Target Production", "Achievement %", "Quality Score"]
    kpi_values = [
        f"{format_number(prod.get('totalCoalProduction', 0.0))} MT",
        f"{format_number(prod.get('totalTargetProduction', 0.0))} MT",
        format_pct(prod.get('productionAchievementPct', prod.get('productionAchievement', 0.0))),
        f"{format_number(val.get('averageValidationScore', 0.0), 1)} / 100"
    ]
    kpi_subtexts = [
        f"Docs: {docs.get('totalDocuments', 0)}",
        f"Variance: {format_number(prod.get('targetVariance', 0.0))} MT",
        f"Valid: {val.get('validDocuments', 0)} | Warn: {val.get('warningDocuments', 0)}",
        f"Rating: {val.get('qualityRating', 'Good')}"
    ]

    for c_idx in range(4):
        # Row 0: label
        cell_lbl = table_kpi.cell(0, c_idx)
        _set_cell_background(cell_lbl, "F1F5F9")
        p0 = cell_lbl.paragraphs[0]
        r0 = p0.add_run(kpi_headers[c_idx])
        r0.font.size = Pt(8.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(71, 85, 105)

        # Row 1: big value
        cell_val = table_kpi.cell(1, c_idx)
        _set_cell_background(cell_val, "F8FAFC")
        p1 = cell_val.paragraphs[0]
        r1 = p1.add_run(kpi_values[c_idx])
        r1.font.size = Pt(13)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(15, 23, 42)

        # Row 2: subtext
        cell_sub = table_kpi.cell(2, c_idx)
        _set_cell_background(cell_sub, "F8FAFC")
        p2 = cell_sub.paragraphs[0]
        r2 = p2.add_run(kpi_subtexts[c_idx])
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 2. Subsidiary Performance
    if subs:
        h_sec2 = doc.add_heading("2. Subsidiary Performance & Production Leaderboard", level=2)
        h_sec2.paragraph_format.space_before = Pt(14)
        h_sec2.paragraph_format.space_after = Pt(6)

        sub_table = doc.add_table(rows=1, cols=6)
        sub_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Rank", "Subsidiary", "Documents", "Total Prod (MT)", "Avg / Record (MT)", "National Share %"]
        for i, h in enumerate(headers):
            cell = sub_table.cell(0, i)
            _set_cell_background(cell, "E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(h)
            r.font.bold = True
            r.font.size = Pt(8.5)

        for s in subs[:10]:
            row = sub_table.add_row()
            vals = [
                f"#{s.get('rank', '-')}",
                str(s.get('subsidiary', 'N/A')),
                str(s.get('documents', 0)),
                format_number(s.get('production', 0.0)),
                format_number(s.get('averageProduction', 0.0)),
                format_pct(s.get('contributionPct', 0.0))
            ]
            for col_idx, val in enumerate(vals):
                cell = row.cells[col_idx]
                p = cell.paragraphs[0]
                r = p.add_run(val)
                r.font.size = Pt(8.5)
                if col_idx in [0, 1, 3]:
                    r.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 3. Validation Audit (VAL001 - VAL010)
    if report_type in ["validation", "executive", "custom"]:
        h_sec3 = doc.add_heading("3. Deterministic Validation Audit (VAL001 - VAL010)", level=2)
        h_sec3.paragraph_format.space_before = Pt(14)
        h_sec3.paragraph_format.space_after = Pt(6)

        v_table = doc.add_table(rows=1, cols=4)
        v_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        v_headers = ["Rule ID", "Rule Name & Description", "Severity", "Occurrences"]
        for i, h in enumerate(v_headers):
            cell = v_table.cell(0, i)
            _set_cell_background(cell, "E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(h)
            r.font.bold = True
            r.font.size = Pt(8.5)

        for rv in rule_violations:
            row = v_table.add_row()
            c0 = row.cells[0].paragraphs[0].add_run(rv["ruleId"])
            c0.font.bold = True
            c0.font.size = Pt(8.5)

            p1 = row.cells[1].paragraphs[0]
            r1_name = p1.add_run(f"{rv['name']}: ")
            r1_name.font.bold = True
            r1_name.font.size = Pt(8.5)
            r1_desc = p1.add_run(rv["description"])
            r1_desc.font.size = Pt(8)

            c2 = row.cells[2].paragraphs[0].add_run(rv["severity"])
            c2.font.bold = True
            c2.font.size = Pt(8.5)
            if rv["severity"] == "Error":
                c2.font.color.rgb = RGBColor(220, 38, 38)
            else:
                c2.font.color.rgb = RGBColor(217, 119, 6)

            c3 = row.cells[3].paragraphs[0].add_run(str(rv["occurrences"]))
            c3.font.bold = True
            c3.font.size = Pt(8.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 4. Mine Performance Table
    if report_type in ["mine_performance", "production", "executive"] and mines:
        h_sec4 = doc.add_heading("4. Mine Performance & Extraction Register", level=2)
        h_sec4.paragraph_format.space_before = Pt(14)
        h_sec4.paragraph_format.space_after = Pt(6)

        m_table = doc.add_table(rows=1, cols=6)
        m_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        m_headers = ["Mine Name", "Subsidiary", "District / State", "Mine Type", "Production (MT)", "Status"]
        for i, h in enumerate(m_headers):
            cell = m_table.cell(0, i)
            _set_cell_background(cell, "E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(h)
            r.font.bold = True
            r.font.size = Pt(8.5)

        for m in mines[:20]:
            row = m_table.add_row()
            vals = [
                m.get("mineName", "N/A"),
                m.get("subsidiary", "N/A"),
                f"{m.get('district', 'N/A')}, {m.get('state', 'N/A')}",
                m.get("mineType", "N/A"),
                format_number(m.get("coalProduction", 0.0)),
                m.get("validationStatus", "Valid")
            ]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                p = cell.paragraphs[0]
                r = p.add_run(str(val))
                r.font.size = Pt(8.5)
                if c_idx == 0:
                    r.font.bold = True
                elif c_idx == 5:
                    r.font.bold = True
                    if val == "Valid":
                        r.font.color.rgb = RGBColor(22, 163, 74)
                    elif val == "Warning":
                        r.font.color.rgb = RGBColor(217, 119, 6)
                    else:
                        r.font.color.rgb = RGBColor(220, 38, 38)

    doc.save(output_path)
    return output_path
