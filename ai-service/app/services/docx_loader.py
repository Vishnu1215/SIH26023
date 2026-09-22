import os
import logging
from docx import Document

logger = logging.getLogger("ai_service.docx_loader")


def load_docx_text(file_path: str) -> dict:
    """
    Extract raw text from Word documents (.docx).
    Reads every paragraph and table text, merging them into structured text.
    Returns: {"text": str, "pages": int, "confidence": None, "language": str, "loaderUsed": "DOCX"}
    """
    if not os.path.exists(file_path):
        err = FileNotFoundError(f"DOCX file not found: {file_path}")
        err.error_code = "CORRUPTED_DOCUMENT"
        raise err

    logger.info(f"Reading DOCX document: {file_path}")
    try:
        doc = Document(file_path)
    except Exception as exc:
        logger.error(f"Failed to parse DOCX file: {exc}")
        err = RuntimeError(f"Corrupted or invalid DOCX document: {exc}")
        err.error_code = "CORRUPTED_DOCUMENT"
        raise err

    text_parts = []

    # Read paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text.strip())

    # Also read tables if present in DOCX
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)

    combined_text = "\n\n".join(text_parts).strip()
    page_estimate = max(1, len(doc.paragraphs) // 20)

    return {
        "text": combined_text,
        "pages": page_estimate,
        "confidence": None,
        "language": "eng",
        "loaderUsed": "DOCX"
    }

