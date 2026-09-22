import os
import logging
import openpyxl

logger = logging.getLogger("ai_service.excel_loader")


def load_excel_text(file_path: str) -> dict:
    """
    Extract raw text from Excel workbooks (.xlsx, .xlsm).
    Reads every worksheet and formats cell rows into readable text.
    Returns: {"text": str, "pages": int, "confidence": None, "language": str, "loaderUsed": "XLSX"}
    """
    if not os.path.exists(file_path):
        err = FileNotFoundError(f"Excel file not found: {file_path}")
        err.error_code = "CORRUPTED_DOCUMENT"
        raise err

    logger.info(f"Reading Excel workbook: {file_path}")
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception as exc:
        logger.error(f"Failed to open Excel workbook: {exc}")
        err = RuntimeError(f"Corrupted or password-protected Excel workbook: {exc}")
        err.error_code = "PASSWORD_PROTECTED" if "password" in str(exc).lower() else "CORRUPTED_DOCUMENT"
        raise err

    sheet_texts = []
    total_sheets = len(workbook.sheetnames)

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        row_lines = []

        for row in sheet.iter_rows(values_only=True):
            # Filter out empty rows
            non_empty_cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip() != ""]
            if non_empty_cells:
                row_lines.append(" | ".join(non_empty_cells))

        if row_lines:
            sheet_content = f"--- Worksheet: {sheet_name} ---\n" + "\n".join(row_lines)
            sheet_texts.append(sheet_content)

    workbook.close()
    combined_text = "\n\n".join(sheet_texts).strip()

    return {
        "text": combined_text,
        "pages": max(1, total_sheets),
        "confidence": None,
        "language": "eng",
        "loaderUsed": "XLSX"
    }

