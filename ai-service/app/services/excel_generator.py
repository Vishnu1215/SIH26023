"""
Excel (.xlsx) Report Generator for Phase 8 using openpyxl.
Produces multi-sheet workbooks with official Ministry branding,
formatted headers, auto-fitted columns, number formatting, and formula totals.
"""

import os
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.services.report_utils import BRANDING, format_number, format_pct, format_datetime, clean_text


def _apply_header_style(cell, text: str):
    """Apply standard Navy header style to Excel cell."""
    cell.value = text
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_subhead_style(cell, text: str):
    """Apply slate subheader style."""
    cell.value = text
    cell.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _autofit_columns(ws, min_width=12, max_width=50):
    """Auto-adjust worksheet column widths based on contents."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, max_width), min_width)


def generate_excel_report(context: Dict[str, Any], output_path: str) -> str:
    """Generate multi-tab Excel report workbook and save to output_path."""
    wb = openpyxl.Workbook()

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    dashboard = context.get("dashboard", {})
    prod = dashboard.get("production", {})
    val = dashboard.get("validation", {})
    docs = dashboard.get("documents", {})
    subs = dashboard.get("subsidiaries", [])
    records = context.get("records", [])
    mines = context.get("minesList", [])
    rule_violations = context.get("ruleViolations", [])

    # -------------------------------------------------------------
    # Sheet 1: Executive Overview
    # -------------------------------------------------------------
    ws_exec = wb.active
    ws_exec.title = "Executive Overview"
    ws_exec.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_exec.merge_cells("A1:E1")
    t1 = ws_exec["A1"]
    t1.value = f"{BRANDING['republic']}  •  {BRANDING['ministry']}"
    t1.font = Font(name="Calibri", size=11, bold=True, color="64748B")
    t1.alignment = Alignment(horizontal="left")

    ws_exec.merge_cells("A2:E2")
    t2 = ws_exec["A2"]
    t2.value = f"{BRANDING['institution']} - Executive Dashboard Summary"
    t2.font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")

    ws_exec["A4"].value = "Generated On:"
    ws_exec["B4"].value = context.get("formattedDate")
    ws_exec["A5"].value = "Analytics Version:"
    ws_exec["B5"].value = f"v{context.get('sourceAnalyticsVersion', 1)}"
    ws_exec["A6"].value = "Records Analyzed:"
    ws_exec["B6"].value = len(records)

    for r in range(4, 7):
        ws_exec[f"A{r}"].font = Font(bold=True, color="475569")

    # KPI Table
    ws_exec.merge_cells("A8:E8")
    _apply_header_style(ws_exec["A8"], "KEY PRODUCTION & COMPLIANCE INDICATORS")

    kpi_rows = [
        ("Total Coal Production (MT)", prod.get("totalCoalProduction", 0.0), "#,##0.00"),
        ("Target Production (MT)", prod.get("totalTargetProduction", 0.0), "#,##0.00"),
        ("Target Variance (MT)", prod.get("targetVariance", 0.0), "#,##0.00"),
        ("Production Achievement %", (prod.get("productionAchievementPct", prod.get("productionAchievement", 0.0))) / 100.0, "0.0%"),
        ("Total Ingested Documents", docs.get("totalDocuments", 0), "#,##0"),
        ("Validated Documents", val.get("validatedDocuments", 0), "#,##0"),
        ("Valid Documents (100% Pass)", val.get("validDocuments", 0), "#,##0"),
        ("Warning Documents", val.get("warningDocuments", 0), "#,##0"),
        ("Error Documents", val.get("errorDocuments", 0), "#,##0"),
        ("Average Validation Quality Score", val.get("averageValidationScore", 0.0), "0.0"),
        ("Quality Rating", val.get("qualityRating", "Good"), "@"),
        ("Missing Fields %", (dashboard.get("quality", {}).get("missingFieldsPercentage", 0.0)) / 100.0, "0.0%")
    ]

    for idx, (label, val_item, num_fmt) in enumerate(kpi_rows, start=9):
        cell_lbl = ws_exec.cell(row=idx, column=1, value=label)
        cell_lbl.font = Font(bold=True, color="1E293B")
        cell_lbl.border = thin_border
        
        cell_val = ws_exec.cell(row=idx, column=2, value=val_item)
        cell_val.number_format = num_fmt
        cell_val.font = Font(name="Calibri", size=11, bold=True, color="0F172A")
        cell_val.border = thin_border

    _autofit_columns(ws_exec)

    # -------------------------------------------------------------
    # Sheet 2: Subsidiary Production
    # -------------------------------------------------------------
    ws_subs = wb.create_sheet(title="Subsidiary Production")
    ws_subs.views.sheetView[0].showGridLines = True

    sub_headers = ["Rank", "Subsidiary", "Documents", "Production (MT)", "Avg / Record (MT)", "National Share %"]
    for c_idx, h in enumerate(sub_headers, start=1):
        _apply_header_style(ws_subs.cell(row=1, column=c_idx), h)

    current_row = 2
    for s in subs:
        ws_subs.cell(row=current_row, column=1, value=s.get("rank", current_row-1)).alignment = Alignment(horizontal="center")
        ws_subs.cell(row=current_row, column=2, value=s.get("subsidiary", "N/A")).font = Font(bold=True)
        ws_subs.cell(row=current_row, column=3, value=s.get("documents", 0)).number_format = "#,##0"
        
        c4 = ws_subs.cell(row=current_row, column=4, value=s.get("production", 0.0))
        c4.number_format = "#,##0.00"
        c4.font = Font(bold=True)

        c5 = ws_subs.cell(row=current_row, column=5, value=s.get("averageProduction", 0.0))
        c5.number_format = "#,##0.00"

        c6 = ws_subs.cell(row=current_row, column=6, value=(s.get("contributionPct", 0.0)) / 100.0)
        c6.number_format = "0.0%"

        for c in range(1, 7):
            ws_subs.cell(row=current_row, column=c).border = thin_border
        current_row += 1

    # Formula Total Row
    if len(subs) > 0:
        ws_subs.cell(row=current_row, column=2, value="TOTAL").font = Font(bold=True)
        c_tot_doc = ws_subs.cell(row=current_row, column=3, value=f"=SUM(C2:C{current_row-1})")
        c_tot_doc.font = Font(bold=True)
        c_tot_doc.number_format = "#,##0"

        c_tot_prod = ws_subs.cell(row=current_row, column=4, value=f"=SUM(D2:D{current_row-1})")
        c_tot_prod.font = Font(bold=True)
        c_tot_prod.number_format = "#,##0.00"

        for c in range(1, 7):
            ws_subs.cell(row=current_row, column=c).border = thin_border
            ws_subs.cell(row=current_row, column=c).fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    _autofit_columns(ws_subs)

    # -------------------------------------------------------------
    # Sheet 3: Mine Performance
    # -------------------------------------------------------------
    ws_mines = wb.create_sheet(title="Mine Performance")
    ws_mines.views.sheetView[0].showGridLines = True

    mine_headers = ["Mine Name", "Subsidiary", "District", "State", "Mine Type", "Production (MT)", "Target (MT)", "Status"]
    for c_idx, h in enumerate(mine_headers, start=1):
        _apply_header_style(ws_mines.cell(row=1, column=c_idx), h)

    m_row = 2
    for m in mines:
        ws_mines.cell(row=m_row, column=1, value=m.get("mineName", "N/A")).font = Font(bold=True)
        ws_mines.cell(row=m_row, column=2, value=m.get("subsidiary", "N/A"))
        ws_mines.cell(row=m_row, column=3, value=m.get("district", "N/A"))
        ws_mines.cell(row=m_row, column=4, value=m.get("state", "N/A"))
        ws_mines.cell(row=m_row, column=5, value=m.get("mineType", "N/A"))
        
        c_prod = ws_mines.cell(row=m_row, column=6, value=m.get("coalProduction", 0.0))
        c_prod.number_format = "#,##0.00"
        c_prod.font = Font(bold=True)

        c_target = ws_mines.cell(row=m_row, column=7, value=m.get("targetProduction", 0.0))
        c_target.number_format = "#,##0.00"

        c_stat = ws_mines.cell(row=m_row, column=8, value=m.get("validationStatus", "Valid"))
        c_stat.alignment = Alignment(horizontal="center")
        if m.get("validationStatus") == "Valid":
            c_stat.font = Font(color="15803D", bold=True)
        elif m.get("validationStatus") == "Warning":
            c_stat.font = Font(color="B45309", bold=True)
        else:
            c_stat.font = Font(color="B91C1C", bold=True)

        for c in range(1, 9):
            ws_mines.cell(row=m_row, column=c).border = thin_border
        m_row += 1

    _autofit_columns(ws_mines)

    # -------------------------------------------------------------
    # Sheet 4: Validation Rule Audit (VAL001 - VAL010)
    # -------------------------------------------------------------
    ws_val = wb.create_sheet(title="Validation Audit")
    ws_val.views.sheetView[0].showGridLines = True

    val_headers = ["Rule ID", "Rule Name", "Severity", "Rule Description", "Violations Detected"]
    for c_idx, h in enumerate(val_headers, start=1):
        _apply_header_style(ws_val.cell(row=1, column=c_idx), h)

    v_row = 2
    for rv in rule_violations:
        c1 = ws_val.cell(row=v_row, column=1, value=rv["ruleId"])
        c1.font = Font(bold=True)
        c1.alignment = Alignment(horizontal="center")
        
        ws_val.cell(row=v_row, column=2, value=rv["name"]).font = Font(bold=True)
        
        c3 = ws_val.cell(row=v_row, column=3, value=rv["severity"])
        c3.alignment = Alignment(horizontal="center")
        c3.font = Font(color="DC2626" if rv["severity"] == "Error" else "D97706", bold=True)

        ws_val.cell(row=v_row, column=4, value=rv["description"])
        
        c5 = ws_val.cell(row=v_row, column=5, value=rv["occurrences"])
        c5.font = Font(bold=True)
        c5.alignment = Alignment(horizontal="right")

        for c in range(1, 6):
            ws_val.cell(row=v_row, column=c).border = thin_border
        v_row += 1

    _autofit_columns(ws_val)

    # Save to disk
    wb.save(output_path)
    return output_path
